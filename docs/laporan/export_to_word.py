import os
import re
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_vertical_align(cell, align='center'):
    tcPr = cell._tc.get_or_add_tcPr()
    vAlign = parse_xml(f'<w:vAlign {nsdecls("w")} w:val="{align}"/>')
    tcPr.append(vAlign)

def set_cell_border(cell, top="000000", bottom="000000", left="000000", right="000000", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{top}"/>\n'
        f'  <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{left}"/>\n'
        f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{bottom}"/>\n'
        f'  <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{right}"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

def set_table_styling(table, num_cols):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
        f'  <w:left w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
        f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

    total_w = 5.52
    if num_cols == 2:
        col_widths = [1.80, 3.72]
    elif num_cols == 3:
        col_widths = [1.50, 2.01, 2.01]
    elif num_cols == 4:
        col_widths = [1.38, 1.38, 1.38, 1.38]
    elif num_cols == 5:
        col_widths = [1.10, 1.10, 1.10, 1.11, 1.11]
    elif num_cols == 6:
        col_widths = [0.92] * 6
    else:
        first_w = max(1.0, total_w * 0.20)
        rem_w = (total_w - first_w) / max(1, num_cols - 1)
        col_widths = [first_w] + [rem_w] * (num_cols - 1)

    for i, col in enumerate(table.columns):
        if i < len(col_widths):
            col.width = Inches(col_widths[i])

    tblGrid = OxmlElement('w:tblGrid')
    for w in col_widths:
        gridCol = OxmlElement('w:gridCol')
        gridCol.set(qn('w:w'), str(int(w * 1440)))
        tblGrid.append(gridCol)
    tblPr.append(tblGrid)

    for r_idx, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if r_idx == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        for idx, cell in enumerate(row.cells):
            if idx < len(col_widths):
                cell.width = Inches(col_widths[idx])
            set_cell_vertical_align(cell, 'center')

def clean_math(math_str):
    s = math_str.strip('$').strip()
    s = s.replace(r'\Delta T', 'ΔT').replace(r'\Delta RH', 'ΔRH').replace(r'\Delta', 'Δ')
    s = re.sub(r'\^\\circ\\text\{C\}', '°C', s)
    s = re.sub(r'\^\\circ C', '°C', s)
    s = re.sub(r'\^\\circ', '°', s)
    s = re.sub(r'\\text\{([^}]+)\}', r'\1', s)
    s = s.replace(r'\times', '×').replace(r'\approx', '≈').replace(r'\pm', '±')
    s = s.replace(r'\le', '≤').replace(r'\ge', '≥')
    s = s.replace(r'\gamma', 'γ').replace(r'\mu', 'μ').replace(r'\sigma', 'σ')
    s = s.replace(r'\implies', ' ⟹ ')
    return s

def clean_markdown_links(text):
    def repl(m):
        label, url = m.groups()
        if label == url or label.startswith('http'):
            return url
        return f"{label} ({url})"
    return re.sub(r'\[([^\]]+)\]\((https?://[^)]+)\)', repl, text)

def format_inline_markdown(paragraph, text, base_font_name='Times New Roman', base_font_size=12, base_color=RGBColor(0,0,0), force_bold=False):
    # Pre-clean escaped characters and links
    text = text.replace(r'\_', '_').replace(r'\*', '*')
    text = clean_markdown_links(text)

    # Robust regex tokenization for code, math, bold, italic
    # Order matters: `code`, $$math$$, $math$, **bold**, *italic*
    pattern = re.compile(r'(`[^`]+`|\$\$.*?\$\$|\$.*?\$|\*\*[^*]+?\*\*|\*[^*\n]+?\*)')
    tokens = pattern.split(text)

    for token in tokens:
        if not token:
            continue
        run = paragraph.add_run()
        run.font.name = base_font_name
        run.font.size = Pt(base_font_size)
        run.font.color.rgb = base_color

        if token.startswith('`') and token.endswith('`'):
            run.text = token[1:-1]
            run.font.name = 'Consolas'
            run.font.size = Pt(base_font_size - 1.5)
            run.font.color.rgb = RGBColor(15, 23, 42)
            if force_bold:
                run.bold = True
        elif (token.startswith('$$') and token.endswith('$$')) or (token.startswith('$') and token.endswith('$')):
            run.text = clean_math(token)
            run.italic = True
            run.font.name = 'Cambria Math'
        elif token.startswith('**') and token.endswith('**'):
            run.text = token[2:-2]
            run.bold = True
        elif token.startswith('*') and token.endswith('*'):
            run.text = token[1:-1]
            run.italic = True
            if force_bold:
                run.bold = True
        else:
            run.text = token
            if force_bold:
                run.bold = True

def add_styled_paragraph(doc, text, space_after=6, line_spacing=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, left_indent=0.0, first_indent=0.0, prefix=None, prefix_bold=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if left_indent > 0:
        p.paragraph_format.left_indent = Inches(left_indent)
    if first_indent != 0.0:
        p.paragraph_format.first_line_indent = Inches(first_indent)

    p.alignment = align
    if align == WD_ALIGN_PARAGRAPH.JUSTIFY:
        pPr = p._p.get_or_add_pPr()
        jc = pPr.find(qn('w:jc'))
        if jc is None:
            pPr.append(parse_xml(f'<w:jc {nsdecls("w")} w:val="both"/>'))
        else:
            jc.set(qn('w:val'), 'both')

    if prefix:
        r_pre = p.add_run(prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(12)
        r_pre.bold = prefix_bold

    format_inline_markdown(p, text, base_font_name='Times New Roman', base_font_size=12)
    return p

def build_cover_page(doc, logo_path):
    p_type = doc.add_paragraph()
    p_type.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_type.paragraph_format.space_before = Pt(0)
    p_type.paragraph_format.space_after = Pt(4)
    p_type.paragraph_format.line_spacing = 1.15
    r = p_type.add_run("LAPORAN RISET PENELITIAN IOT AGROMETEOROLOGI")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(24)
    p_title.paragraph_format.line_spacing = 1.15
    r = p_title.add_run("RANCANG BANGUN SISTEM STASIUN CUACA AGROMETEOROLOGI BERBASIS MULTI-SENSOR DAN DUAL-CONTROLLER TERINTEGRASI REST API LOKAL SERTA CLOUD FIRESTORE DENGAN MODEL EVAPOTRANSPIRASI FAO-56 PENMAN-MONTEITH")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True

    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(12)
        p_logo.paragraph_format.space_after = Pt(28)
        r_logo = p_logo.add_run()
        r_logo.add_picture(logo_path, width=Inches(2.1))

    p_disusun = doc.add_paragraph()
    p_disusun.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_disusun.paragraph_format.space_before = Pt(8)
    p_disusun.paragraph_format.space_after = Pt(4)
    r = p_disusun.add_run("Disusun Oleh:")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.italic = True

    p_nama = doc.add_paragraph()
    p_nama.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_nama.paragraph_format.space_before = Pt(2)
    p_nama.paragraph_format.space_after = Pt(2)
    r = p_nama.add_run("HAIKAL JIBRAN A.")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.bold = True

    p_nim = doc.add_paragraph()
    p_nim.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_nim.paragraph_format.space_before = Pt(0)
    p_nim.paragraph_format.space_after = Pt(2)
    r = p_nim.add_run("NIM: 20240810005")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.bold = True

    p_kelas = doc.add_paragraph()
    p_kelas.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_kelas.paragraph_format.space_before = Pt(0)
    p_kelas.paragraph_format.space_after = Pt(40)
    r = p_kelas.add_run("Kelas: TINFC-2024-01")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)

    p_prodi = doc.add_paragraph()
    p_prodi.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_prodi.paragraph_format.space_before = Pt(0)
    p_prodi.paragraph_format.space_after = Pt(2)
    r = p_prodi.add_run("PROGRAM STUDI S1 TEKNIK INFORMATIKA")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True

    p_fkom = doc.add_paragraph()
    p_fkom.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_fkom.paragraph_format.space_before = Pt(0)
    p_fkom.paragraph_format.space_after = Pt(2)
    r = p_fkom.add_run("FAKULTAS ILMU KOMPUTER")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True

    p_univ = doc.add_paragraph()
    p_univ.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_univ.paragraph_format.space_before = Pt(0)
    p_univ.paragraph_format.space_after = Pt(2)
    r = p_univ.add_run("UNIVERSITAS KUNINGAN")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True

    p_year = doc.add_paragraph()
    p_year.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_year.paragraph_format.space_before = Pt(0)
    p_year.paragraph_format.space_after = Pt(0)
    r = p_year.add_run("2026")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.bold = True

    doc.add_page_break()

def export_markdown_to_docx(md_path, docx_path, logo_path):
    print(f"Reading markdown source: {md_path}")
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()

    # A4 Page Setup: Margins 4 cm Left, 4 cm Top, 3 cm Right, 3 cm Bottom
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(1.57)    # 4.0 cm
    section.top_margin = Inches(1.57)     # 4.0 cm
    section.right_margin = Inches(1.18)   # 3.0 cm
    section.bottom_margin = Inches(1.18)  # 3.0 cm

    build_cover_page(doc, logo_path)

    # Configure Normal style to default to Justified across the document
    style_normal = doc.styles['Normal']
    style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(12)
    style_normal.font.color.rgb = RGBColor(0, 0, 0)

    pPr_style = style_normal._element.get_or_add_pPr()
    jc_style = pPr_style.find(qn('w:jc'))
    if jc_style is None:
        pPr_style.append(parse_xml(f'<w:jc {nsdecls("w")} w:val="both"/>'))
    else:
        jc_style.set(qn('w:val'), 'both')

    base_dir = os.path.dirname(md_path)

    in_code_block = False
    code_lines = []
    code_lang = ""

    table_lines = []
    in_table = False

    skip_header = True

    i = 0
    while i < len(lines):
        raw_line = lines[i].rstrip('\r\n')
        line = raw_line.strip()

        # Skip markdown cover header text until section 1 starts
        if skip_header:
            if line.startswith('## 1.') or line.startswith('## BAB 1'):
                skip_header = False
            else:
                i += 1
                continue

        # Handle Code Block
        if line.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_lang = line[3:].strip()
                code_lines = []
            else:
                in_code_block = False
                # Insert code block table
                table = doc.add_table(rows=1, cols=1)
                table.autofit = False
                table.columns[0].width = Inches(5.52)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER

                tblPr = table._tbl.tblPr
                tbl_borders = parse_xml(
                    f'<w:tblBorders {nsdecls("w")}>\n'
                    f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>\n'
                    f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>\n'
                    f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>\n'
                    f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>\n'
                    f'  <w:insideH w:val="none"/>\n'
                    f'  <w:insideV w:val="none"/>\n'
                    f'</w:tblBorders>'
                )
                tblPr.append(tbl_borders)

                cell = table.cell(0, 0)
                set_cell_background(cell, "F8FAFC")
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

                code_text = "\n".join(code_lines)
                run = p.add_run(code_text)
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(15, 23, 42)

                # Spacing after code box
                p_spacer = doc.add_paragraph()
                p_spacer.paragraph_format.space_before = Pt(0)
                p_spacer.paragraph_format.space_after = Pt(6)
            i += 1
            continue

        if in_code_block:
            code_lines.append(raw_line)
            i += 1
            continue

        # Handle Tables
        if line.startswith('|') and line.endswith('|'):
            if not in_table:
                in_table = True
                table_lines = [line]
            else:
                table_lines.append(line)
            i += 1
            continue
        elif in_table:
            # End of table, process it
            in_table = False
            valid_rows = []
            for tl in table_lines:
                parts = [p.strip() for p in tl.split('|')[1:-1]]
                if all(re.match(r'^:?-+:?$', p) for p in parts if p):
                    continue
                valid_rows.append(parts)

            if valid_rows:
                num_cols = max(len(r) for r in valid_rows)
                tbl = doc.add_table(rows=len(valid_rows), cols=num_cols)
                set_table_styling(tbl, num_cols)

                for r_idx, r_data in enumerate(valid_rows):
                    for c_idx in range(num_cols):
                        val = r_data[c_idx] if c_idx < len(r_data) else ""
                        c = tbl.cell(r_idx, c_idx)
                        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
                        set_cell_border(c, sz="4")
                        cp = c.paragraphs[0]
                        cp.paragraph_format.space_before = Pt(0)
                        cp.paragraph_format.space_after = Pt(0)
                        cp.paragraph_format.line_spacing = 1.15
                        if r_idx == 0:
                            set_cell_background(c, "F1F5F9")
                            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            format_inline_markdown(cp, val, base_font_size=10, force_bold=True)
                        else:
                            cp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                            format_inline_markdown(cp, val, base_font_size=10)

                p_after_tbl = doc.add_paragraph()
                p_after_tbl.paragraph_format.space_after = Pt(6)

        # Check for images: ![Caption](path)
        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)$', line)
        if img_match:
            caption, img_rel_path = img_match.groups()
            img_rel_path = img_rel_path.lstrip('./')
            img_full_path = os.path.join(base_dir, img_rel_path)
            if not os.path.exists(img_full_path):
                img_full_path = os.path.join(base_dir, 'assets', os.path.basename(img_rel_path))

            if os.path.exists(img_full_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(2)
                r_img = p_img.add_run()
                r_img.add_picture(img_full_path, width=Inches(5.2))

                # Look ahead: if next line is an italic caption like *Gambar X.Y: ...*, use it and skip that line
                next_caption = ""
                if i + 1 < len(lines):
                    next_line = lines[i + 1].strip()
                    if next_line.startswith('*Gambar') or next_line.startswith('*Figura') or next_line.startswith('_Gambar'):
                        next_caption = next_line.strip('*_ ')
                        i += 1  # Skip next line so it is not rendered twice

                final_caption = next_caption if next_caption else caption
                if final_caption:
                    p_cap = doc.add_paragraph()
                    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_cap.paragraph_format.space_before = Pt(2)
                    p_cap.paragraph_format.space_after = Pt(10)
                    r_cap = p_cap.add_run(final_caption)
                    r_cap.font.name = 'Times New Roman'
                    r_cap.font.size = Pt(10)
                    r_cap.italic = True
            else:
                add_styled_paragraph(doc, f"*[Gambar tidak ditemukan: {img_rel_path}]*")
            i += 1
            continue

        # Handle Headings
        if line.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            format_inline_markdown(p, line[5:], base_font_size=12, force_bold=True)
            i += 1
            continue
        elif line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            format_inline_markdown(p, line[4:], base_font_size=12, force_bold=True)
            i += 1
            continue
        elif line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            format_inline_markdown(p, line[3:], base_font_size=13, force_bold=True)
            i += 1
            continue
        elif line.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            format_inline_markdown(p, line[2:], base_font_size=14, force_bold=True)
            i += 1
            continue

        # Horizontal rule ---
        if line == '---':
            i += 1
            continue

        # Empty lines
        if not line:
            i += 1
            continue

        # Sub-list items: indented by 2 or 4 spaces with dash or bullet
        # e.g., "    - **Tingkat 1 - Administrator:** ..."
        sub_list_match = re.match(r'^(?: {2,6}|\t)(?:\*|-)\s+(.*)$', raw_line)
        if sub_list_match:
            item_text = sub_list_match.group(1).strip()
            add_styled_paragraph(
                doc,
                item_text,
                space_after=3,
                line_spacing=1.5,
                align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                left_indent=0.55,
                first_indent=-0.22,
                prefix="–   ",
                prefix_bold=False
            )
            i += 1
            continue

        # Top-level list items: bullet (*, -) or numbered (1., 2.)
        list_match = re.match(r'^(\*|-|(\d+)\.)\s+(.*)$', line)
        if list_match:
            bullet_sym = list_match.group(1)
            num_idx = list_match.group(2)
            item_text = list_match.group(3).strip()

            if num_idx:
                prefix = f"{num_idx}.  "
            else:
                prefix = "•   "

            add_styled_paragraph(
                doc,
                item_text,
                space_after=4,
                line_spacing=1.5,
                align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                left_indent=0.28,
                first_indent=-0.28,
                prefix=prefix,
                prefix_bold=True
            )
            i += 1
            continue

        # Regular body paragraph: strictly JUSTIFIED
        add_styled_paragraph(
            doc,
            line,
            space_after=6,
            line_spacing=1.5,
            align=WD_ALIGN_PARAGRAPH.JUSTIFY,
            left_indent=0.0,
            first_indent=0.0
        )
        i += 1

    print(f"Saving compiled document to: {docx_path}")
    doc.save(docx_path)
    print("Export completed successfully!")

if __name__ == '__main__':
    md_file = "/home/ghiffa/Documents/Projects_IoT/PlatformIO_Workspace/weather-station-agrinex/docs/laporan/DOKUMENTASI_LENGKAP_SISTEM_AGRINEX.md"
    docx_file = "/home/ghiffa/Documents/Projects_IoT/PlatformIO_Workspace/weather-station-agrinex/docs/laporan/DOKUMENTASI_LENGKAP_SISTEM_AGRINEX.docx"
    logo_file = "/home/ghiffa/Documents/Projects_IoT/PlatformIO_Workspace/weather-station-agrinex/docs/laporan/assets/logo_uniku.png"
    export_markdown_to_docx(md_file, docx_file, logo_file)
